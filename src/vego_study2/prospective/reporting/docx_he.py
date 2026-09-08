"""Hebrew right-to-left Word documents: scientific report, executive summary, rater cards.

Uses python-docx (locked ``thesis`` dependency group).  Complex-script sizing
(``w:szCs``) and paragraph direction (``w:bidi``) are set explicitly because
python-docx does not expose them.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .. import constants as c
from .analysis import fmt_counts

HEB = "David"
LATIN = "Calibri"
DARK = RGBColor(0x1F, 0x3A, 0x5F)
BODY = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x59, 0x59, 0x59)


def _rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    paragraph.alignment = align
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    jc = pPr.find(qn("w:jc"))
    if jc is not None:
        jc.addprevious(bidi)
    else:
        pPr.append(bidi)


def _run(paragraph, text: str, size: float = 11, bold: bool = False, italic: bool = False, color: RGBColor = BODY, rtl: bool = True):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = LATIN
    rPr = run._r.get_or_add_rPr()
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rPr.insert(0, fonts)
    fonts.set(qn("w:ascii"), LATIN)
    fonts.set(qn("w:hAnsi"), LATIN)
    fonts.set(qn("w:cs"), HEB)
    if bold:
        rPr.append(OxmlElement("w:bCs"))
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(size * 2)))
    rPr.append(szcs)
    if rtl:
        rPr.append(OxmlElement("w:rtl"))
    return run


def para(doc, text: str = "", size: float = 11, bold: bool = False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after: float = 5,
         color: RGBColor = BODY, italic: bool = False, rtl: bool = True):
    p = doc.add_paragraph()
    if rtl:
        _rtl(p, align)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    if text:
        _run(p, text, size=size, bold=bold, italic=italic, color=color, rtl=rtl)
    return p


def heading(doc, text: str, level: int = 1):
    size = {0: 20, 1: 15, 2: 12.5, 3: 11.5}[level]
    p = para(doc, text, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=4, color=DARK)
    p.paragraph_format.space_before = Pt(10 if level <= 1 else 6)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text: str, size: float = 10.5):
    p = para(doc, "", size=size, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=2)
    p.paragraph_format.right_indent = Cm(0.5)
    _run(p, "• " + text, size=size)
    return p


def table(doc, header: list[str], rows: list[list[Any]], size: float = 9, widths_cm: list[float] | None = None, rtl: bool = True):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    if rtl:
        bidi = OxmlElement("w:bidiVisual")
        tblPr.append(bidi)
    for j, text in enumerate(header):
        cell = tbl.rows[0].cells[j]
        cell.paragraphs[0].text = ""
        _rtl(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
        _run(cell.paragraphs[0], str(text), size=size, bold=True, color=DARK)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), "DDEBF7")
        cell._tc.get_or_add_tcPr().append(shading)
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            cell = tbl.rows[i].cells[j]
            cell.paragraphs[0].text = ""
            _rtl(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
            _run(cell.paragraphs[0], "" if value is None else str(value), size=size)
    if widths_cm:
        for row in tbl.rows:
            for j, width in enumerate(widths_cm):
                row.cells[j].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def figure(doc, png: Path, caption: list[str], width_cm: float = 15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(png), width=Cm(width_cm))
    for index, line in enumerate(caption):
        cp = para(doc, line, size=8.5, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0 if index < len(caption) - 1 else 8,
                  color=GREY, bold=index == 0)
        cp.paragraph_format.keep_with_next = index < len(caption) - 1


def setup(doc, landscape: bool = False) -> None:
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    else:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2.0))
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._r.append(el)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._r.append(instr)
    run.font.size = Pt(9)


def page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _f(value: Any, digits: int = 4) -> str:
    if value is None:
        return "לא זמין"
    if isinstance(value, float):
        return f"{value:.{digits}f}"
    return str(value)


def _s(value: float | None) -> str:
    return "לא זמין" if value is None else f"{value:.0f}"


CAPTION_HE = {
    "C1": ["תרשים C1: שלמות פלט ותקינות מבנית לכל מקרה מזווג",
           "מקור: public-aggregate.json של הריצה · תנאי: ON ו-OFF · מכנה: {n} מקרים מזווגים מתוכננים",
           "סיווג ראיות: {evidence} · מדד: M1 הפקת פריט, M2 תקינות מול החוזה המשותף",
           "מגבלה: קורפוס אחד, מודל אחד, ללא אמת-קרקע; שלמות אינה נכונות"],
    "C2": ["תרשים C2: בקשות לספק ועלות מתומחרת לפי תנאי",
           "מקור: public-aggregate.json · תנאי: ON (מפוצל לרמת-הגדרה ומשויך-למקרה) ו-OFF · מכנה: כל הבקשות כולל ניסיונות חוזרים",
           "סיווג ראיות: {evidence} · מדד: M3 עלות ביצוע מדיווח שימוש במחירון קפוא (0.20 / 1.20 דולר למיליון אסימונים)",
           "מגבלה: עלות היא תיאור של ריצה זו ולא טענת עדיפות; עלות רמת-ההגדרה של ON משותפת לכל המקרים"],
    "C3": ["תרשים C3: זמן שחלף לפי תנאי ולפי מקרה",
           "מקור: public-aggregate.json · תנאי: ON ו-OFF · מכנה: לכל תנאי מדויק; ל-ON לכל מקרה — טווח מהבקשה המשויכת הראשונה לאחרונה תחת מקביליות 2",
           "סיווג ראיות: {evidence} · מדד: M4 שניות שעון-קיר",
           "מגבלה: כולל שונות בזמן התגובה של הספק ביום אחד; טווחי ON לכל מקרה חופפים ומקורבים"],
    "C4": ["תרשים C4: אפיזודות שאלה–תשובה, שאלות ותשובות לכל מקרה (ON בלבד)",
           "מקור: public-aggregate.json · תנאי: ON · מכנה: אפיזודות שנרשמו ושויכו למקרה",
           "סיווג ראיות: {evidence} · מדד: M6 זמינות ראיות תקשורת, M7 מדדי תקשורת",
           "מגבלה: ל-OFF אין אפיזודות מבנית (NOT_AVAILABLE); אפס אפיזודות תחת ON הוא תצפית תקפה"],
    "C5": ["תרשים C5: תוויות מועמדות של Detector-v1 לכל אפיזודה (ON בלבד, דיווח בלבד)",
           "מקור: public-aggregate.json · תנאי: ON; ל-OFF לא רלוונטי · מכנה: {complete} אפיזודות שלמות מדעית",
           "סיווג ראיות: {evidence} · מדד: M8 סיווג Detector-v1 (STRONG = S1 או S3 או S7; WEAK = S2 או S6)",
           "מגבלה: אין תוויות אנושיות; תווית היא מועמדת לבדיקה ולא בעיה מאומתת"],
    "C6": ["תרשים C6: שורות מיפוי ופרגמנטים לא מכוסים לכל מקרה מזווג",
           "מקור: public-aggregate.json · תנאי: ON ו-OFF · מכנה: פריטים שהושלמו (מקרה ללא פריט מוצג כאפס)",
           "סיווג ראיות: {evidence} · מדד: M1/M9 מתארים מבניים של חוזה הפלט המשותף",
           "מגבלה: ספירות הן מתארים מבניים ולא שיפוט איכות; יותר שורות אינו טוב יותר"],
    "C7": ["תרשים C7: הערכה אנושית (סעיף E)",
           "מקור: המחוון בעברית; הכרטיסים בשורש הפרטי · תנאי: אפיזודות ON · מכנה: כרטיסים שנוקדו על ידי שני מעריכים (0 עד כה)",
           "סיווג ראיות: NOT_MEASURED · מדד: M10 החלטות עיוורות REVIEW_WORTHY / NOT_REVIEW_WORTHY / INSUFFICIENT_INFORMATION",
           "מגבלה: דבר אינו מוצג עד שקיימים ניקודים אנושיים אמיתיים; שיפוט של מודל אינו תחליף לאדם"],
}


def _caption(key: str, analysis: dict[str, Any]) -> list[str]:
    return [line.format(n=len(analysis["case_ids"]), evidence=_evidence_he(analysis), complete=analysis["detector"]["scientific_complete"])
            for line in CAPTION_HE[key]]


def _evidence_he(analysis: dict[str, Any]) -> str:
    return {"PROSPECTIVE EMPIRICAL EVIDENCE": "ראיות אמפיריות פרוספקטיביות",
            "ENGINEERING-ONLY FIXTURE": "תוצר הנדסי בלבד (ספק מדומה)"}.get(analysis["evidence_class"], analysis["evidence_class"])


def build_report(analysis: dict[str, Any], manifest: dict[str, Any], selection: dict[str, Any], budget_doc: dict[str, Any],
                 charts: dict[str, Path], cards: dict[str, Any], out_path: Path) -> Path:
    doc = Document()
    setup(doc)
    cond = analysis["conditions"]
    on, off = cond[c.CONDITION_ON], cond[c.CONDITION_OFF]
    det = analysis["detector"]
    diff = analysis["differences"]
    n = len(analysis["case_ids"])
    evidence_he = _evidence_he(analysis)
    fixture = analysis["mode"] != "LIVE"

    para(doc, "VEGO-AI · מחקר 2", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, color=GREY)
    heading(doc, "ניסוי פרוספקטיבי מזווג: VEGO‑AI_ON לעומת VEGO‑AI_OFF", 0)
    para(doc, "דוח מדעי למנחים — השוואה תפעולית תיאורית של זרם עבודה רב-סוכני מול זרם עבודה ישיר מותאם, על קורפוס ציבורי אחד", size=12, align=WD_ALIGN_PARAGRAPH.RIGHT, color=GREY)
    if fixture:
        para(doc, "אזהרה: מסמך זה נבנה מריצת בדיקה עם ספק מדומה. כל המספרים הם תוצר הנדסי בלבד ואינם ראיות אמפיריות.", size=11, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, color=RGBColor(0xA6, 0x29, 0x1F))
    table(doc, ["פריט", "ערך"], [
        ["מזהה ריצה", analysis["run_id"]],
        ["ראש Git של הביצוע", analysis["execution_git_sha"]],
        ["SHA-256 של המניפסט הקפוא", analysis["manifest_sha256"]],
        ["ספק / מודל", f"{analysis['model']['provider']} / {analysis['model']['model']}"],
        ["התחלה / סיום (UTC)", f"{analysis['started_at']} / {analysis['completed_at']}"],
        ["סיווג ראיות", evidence_he],
        ["מעמד המסמך", "טיוטה למנחים; אינו כולל הערכה אנושית (NOT_MEASURED)"],
    ], size=9, widths_cm=[4.5, 12.5])

    heading(doc, "1. תקציר מנהלים", 1)
    para(doc, f"בוצעה ריצה אחת, מאושרת מראש ומוקפאת לפני כל פלט, של הניסוי המזווג על {n} מקרים שנבחרו בכלל בר-ביצוע עם סיד קבוע מתוך {selection['eligible_case_count']} מודלים מועמדים כשירים בקורפוס Text2UML AirTravel, באמצעות המודל {analysis['model']['model']} בלבד ותחת תקרת הוצאה קשיחה של 6.00 דולר. שני התנאים קיבלו אותם מקרים, אותה תיאור-תחום, אותה סכמת פלט, אותה תקרת פלט ואותה מדיניות בקשות; ההבדל היחיד הוא זרם העבודה: ארבעה סוכנים עם שאלה–תשובה מתועדת (ON) לעומת קריאה ישירה אחת למקרה (OFF).")
    bullet(doc, f"שלמות פלט (M1): ON {on['completed']} מתוך {on['planned']} פריטים; OFF {off['completed']} מתוך {off['planned']}. שני התנאים השלימו יחד {diff['both_completed']} מקרים.")
    bullet(doc, f"תקינות מבנית (M2): ON {on['summary_consistent_count']} מתוך {on['completed']} פריטים עם סיכום כיסוי עקבי; OFF {off['summary_consistent_count']} מתוך {off['completed']}.")
    bullet(doc, f"בקשות ועלות (M3): ON {on['requests']} בקשות ({on['setting_level_requests']} ברמת ההגדרה, {on['case_attributed_requests']} משויכות למקרים), {on['cost_usd']:.4f} דולר; OFF {off['requests']} בקשות, {off['cost_usd']:.4f} דולר. סך הכול {analysis['budget']['actual_cost_usd']:.4f} דולר מתוך רזרבציה של {analysis['whole_study_reservation_usd']:.4f} ותקרה של 6.00.")
    bullet(doc, f"זמן (M4): ON {_s(on['elapsed_seconds'])} שניות; OFF {_s(off['elapsed_seconds'])} שניות לכל התנאי.")
    bullet(doc, f"כשלים טכניים (M5): ON {on['technical_failures']}, OFF {off['technical_failures']}; ניסיונות חוזרים על תקלות תעבורה: ON {on['retried_requests']}, OFF {off['retried_requests']}.")
    bullet(doc, f"ראיות תקשורת (M6): ב-ON נרשמו {det['episodes_total']} אפיזודות שאלה–תשובה ({det['scientific_complete']} שלמות מדעית); ב-OFF אין אפיזודות מבנית.")
    bullet(doc, f"Detector-v1 (M8, ON בלבד, דיווח בלבד): {fmt_counts(det['classifications'])}; ב-OFF: לא רלוונטי.")
    bullet(doc, "הערכה אנושית (M10): NOT_MEASURED — עד שני מעריכים בלתי תלויים. אין כאן טענת עדיפות, נכונות התרעות, דיוק או תועלת אנושית.")

    heading(doc, "2. שאלות המחקר וגבול הטענות", 1)
    para(doc, "שאלה ראשית (שני התנאים): כאשר מודל הספק, מקרי הקלט, חומר התחום, סכמת הפלט, תקרת אסימוני הפלט, הטמפרטורה, הזמן הקצוב ומדיניות התקציב קבועים, כיצד נבדלים זרם העבודה הרב-סוכני של VEGO-AI וזרם עבודה ישיר מותאם בשלמות הפלט, בתקינות המבנית, בעלות הביצוע, בזמן, בשיעור הכשלים ובזמינות של ראיות תקשורת ניתנות להסבר?")
    para(doc, "שאלה משנית (ON בלבד): האם Detector-v1 הקפוא מזהה אפיזודות שאלה–תשובה הראויות לבדיקה אנושית בעדיפות? התשובה תלויה בניקוד אנושי עיוור ועד אז היא NOT_MEASURED.")
    para(doc, "זהו ניסוי פרוספקטיבי חקרני מזווג. הוא אינו טענה לעדיפות מערכת, לתועלת אנושית, לנכונות התרעות, לדיוק, לרגישות, ל-F1, לסיבתיות, לייצוגיות או להכללה. לא נרשמה השערה כיוונית; הניתוח תיאורי ומזווג לכל מקרה.", italic=True)

    heading(doc, "3. שיטה", 1)
    heading(doc, "3.1 התנאים", 2)
    table(doc, ["מאפיין", "VEGO-AI_ON", "VEGO-AI_OFF"], [
        ["זרם עבודה", "צינור ארבעת הסוכנים ללא שינוי: יועץ שפה, יועץ תחום, בוחן מודל, חוקר שונות", "בקשה ישירה אחת לכל מקרה"],
        ["פירוק לסוכנים / שאלה–תשובה / לולאת סבבים", "כן / כן, מוגבל ומתועד / MAX_QA_ROUNDS כמשלוח", "לא / לא / אין"],
        ["Detector-v1", "רלוונטי; תווית דיווח בלבד", "לא רלוונטי (אין אפיזודות; לא \"אפס התרעות\")"],
        ["מודל, פרמטרים, סכמת פלט, תקרת פלט, זמן קצוב, ניסיונות חוזרים, תקציב", "זהים", "זהים"],
        ["טמפרטורה", "לא נשלחת (ברירת מחדל של הספק)", "לא נשלחת (זהה מבנית)"],
    ], size=8.5, widths_cm=[4.2, 6.6, 6.2])
    para(doc, "OFF אינו \"בלי בינה מלאכותית\". שני התנאים משתמשים באותה מחלקת לקוח מוגנת, כך שצורת הבקשה, ניסיונות הפענוח החוזרים והטמפרטורה זהים מבנית ולא רק מוצהרים.", size=10)

    heading(doc, "3.2 נתונים ובחירת מקרים", 2)
    para(doc, f"הקורפוס הוא Text2UML AirTravel (מזהה {selection['corpus_id']}, קומיט {selection['corpus_commit'][:12]}), ציבורי, חיצוני ונוצר על ידי מודלי שפה; אינו נתוני סטודנטים, אינו Cheers או ParkWise ואינו ייצוגי. המסגרת הכשירה כוללת {selection['eligible_case_count']} מודלים מועמדים ותיאור תחום אחד, וכל תקציר SHA-256 אומת מחדש לפני כל בקשה.")
    para(doc, f"כלל הבחירה בר-הביצוע: {selection['selection_rule'].split(';')[0]} עם סיד {selection['selection_seed']}. נבחרו: {', '.join(selection['selected_case_ids'])}. הוצאו (לא נדגמו): {', '.join(selection['excluded_case_ids'])}. שלושה קבצים נבחרים זהים בייט-לבייט לקלטים של מחקר 1 ({', '.join(f'{k}→{v}' for k, v in selection['study1_overlap_by_digest'].items())}); שום פלט של מחקר 1 אינו בשימוש חוזר.", size=10)

    heading(doc, "3.3 תקציב ובחירת פרוטוקול", 2)
    menu = budget_doc["menu"]
    para(doc, f"מחירון קפוא: 0.20 / 1.20 דולר למיליון אסימוני קלט / פלט. רזרבציה לכל בקשה = (12,000 × 0.20 + 16,384 × 1.20) / 10⁶ = {menu['per_request_reserve_usd']:.7f} דולר. תקרה קשיחה 6.00; תקרת שמירה 5.90 (שולי בטיחות של 0.10 לחריגת אסימוני קלט מעל הרזרבה בבקשה האחרונה). תקרת בקשות כוללת = ⌊5.90 / {menu['per_request_reserve_usd']:.7f}⌋ = {menu['total_request_cap']}.", size=10)
    table(doc, ["אפשרות", "מקרים נפרדים", "חזרות", "יחידות מזווגות", "בקשות שמורות", "רזרבציה (דולר)", "מתאים"], [
        [o["option"], o["distinct_cases"], o["repeats_per_case"], o["paired_units"], o["requests_reserved"], f"{o['reservation_usd']:.4f}", "כן" if o["fits_guard_ceiling"] else "לא"]
        for o in menu["options"]
    ], size=8.5)
    para(doc, "נבחרה אפשרות A (כיסוי מרבי): השאלה הראשית היא השוואה מזווגת לכל מקרה, ולכן יחידות מזווגות נפרדות שוות יותר מחזרות של מערכת סטוכסטית בטמפרטורת ברירת מחדל; אפשרות B מחייבת אותה רזרבציה עם מחצית המקרים הנפרדים, ואפשרות C משאירה שני שלישים מהמדגם שניתן לממן ללא שימוש.", size=10)
    para(doc, f"תקרות תפעוליות: ON {manifest['caps']['on_requests_per_case']} בקשות למקרה ({manifest['caps'][c.CONDITION_ON]} בסך הכול), OFF {manifest['caps']['off_requests_per_case']} למקרה ({manifest['caps'][c.CONDITION_OFF]}). החסם התאורטי של ON (82 + 61N קריאות לוגיות בעשרה סבבים, עד שלושה ניסיונות פענוח לכל קריאה) חורג מהתקרה בכל N ≥ 4 תחת רזרבציית פלט מלאה; לכן מנגנון האכיפה הוא שומר הרזרבציה לכל בקשה, והתקרות התפעוליות מכריעות על הקבלה. מחקר 1 (43 בקשות לארבעה מקרים, 0.135 דולר) שימש לכיול בלבד ולא כחסם עליון.", size=10)

    heading(doc, "3.4 פרמטרים קפואים ושערים", 2)
    passed = sum(1 for g in analysis["gates"].values() if g["passed"])
    table(doc, ["פריט", "ערך"], [
        ["פרמטרי בקשה", f"model, messages [system, user], max_completion_tokens {analysis['model']['request_parameters']['max_completion_tokens']}; temperature/seed/response_format/tools לא נשלחים"],
        ["יעד רשת מותר", ", ".join(analysis["model"]["allowed_hosts"])],
        ["אישורים", "OPENAI_API_KEY נקרא רק על ידי ה-SDK של הספק; המערכת אינה מדפיסה, רושמת, מגבבת או משדרת את הערך"],
        ["סדר ביצוע / מקביליות", "OFF ואז ON / 2 מקרים"],
        ["זמן קצוב", f"בקשה 180 שניות; OFF {manifest['run_timeouts_seconds'][c.CONDITION_OFF]} שניות; ON {manifest['run_timeouts_seconds'][c.CONDITION_ON]} שניות"],
        ["ניסיונות חוזרים", "ניסיון תעבורה חוזר אחד לכל בקשה; ניסיונות פענוח חוזרים של הלקוח המוגן זהים לשני התנאים; כל ניסיון נספר בתקרות ובתקציב"],
    ], size=8.5, widths_cm=[4.2, 12.8])
    para(doc, f"שערים לפני הביצוע ({passed} מתוך {len(analysis['gates'])} עברו; כל שער שנכשל מסרב לבצע):", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    table(doc, ["שער", "עבר", "פירוט"], [
        [name, "כן" if gate["passed"] else "לא", gate["detail"]] for name, gate in analysis["gates"].items()
    ], size=8, widths_cm=[5.2, 1.3, 10.5])

    page_break(doc)
    heading(doc, "4. תוצאות A — מדדים תפעוליים משותפים", 1)
    table(doc, ["מדד", "VEGO-AI_ON", "VEGO-AI_OFF"], [
        ["פריטים שהושלמו / מתוכננים (M1)", f"{on['completed']} / {on['planned']}", f"{off['completed']} / {off['planned']}"],
        ["פריטים לא תקינים סכמטית / כשלים טכניים (M2, M5)", f"{on['schema_invalid']} / {on['technical_failures']}", f"{off['schema_invalid']} / {off['technical_failures']}"],
        ["סיכום כיסוי עקבי מתוך שהושלמו (M2)", f"{on['summary_consistent_count']} / {on['completed']}", f"{off['summary_consistent_count']} / {off['completed']}"],
        ["בקשות (כולל ניסיונות חוזרים) (M3)", f"{on['requests']} (רמת הגדרה {on['setting_level_requests']}, משויך {on['case_attributed_requests']})", str(off["requests"])],
        ["אסימוני קלט / פלט (M3)", f"{on['prompt_tokens']} / {on['completion_tokens']}", f"{off['prompt_tokens']} / {off['completion_tokens']}"],
        ["עלות בדולר (M3)", f"{on['cost_usd']:.6f}", f"{off['cost_usd']:.6f}"],
        ["עלות לפריט שהושלם (M9)", _f(on["cost_per_completed_artifact_usd"], 6), _f(off["cost_per_completed_artifact_usd"], 6)],
        ["בקשות לפריט שהושלם (M9)", _f(on["requests_per_completed_artifact"], 2), _f(off["requests_per_completed_artifact"], 2)],
        ["זמן שחלף לכל התנאי, שניות (M4)", _s(on["elapsed_seconds"]), _s(off["elapsed_seconds"])],
        ["תקלות תעבורה / ניסיונות חוזרים (M5)", f"{on['transport_errors']} / {on['retried_requests']}", f"{off['transport_errors']} / {off['retried_requests']}"],
        ["אפיזודות שאלה–תשובה (M6)", str(det["episodes_total"]), "אין מבנית (NOT_AVAILABLE)"],
        ["Detector-v1", "רלוונטי", "לא רלוונטי"],
    ], size=8.5, widths_cm=[6.2, 5.6, 5.2])
    acc = analysis["accounting"]
    if acc["unrecorded_in_flight_requests"]:
        para(doc, f"הערת חשבונאות: השומר ספר {acc['receipt_requests']} בקשות שהונפקו, אך ביומן הקריאות {acc['ledger_requests']} שורות. {acc['unrecorded_in_flight_requests']} בקשה/ות שהיו בטיסה ברגע שהתקרה עצרה את התנאי בוטלו ולא החזירו נתוני שימוש; הן נספרות כבקשות שהונפקו ומחויבות ברזרבציה המלאה, ולכן חסם עליון להוצאה בפועל הוא {acc['spend_upper_bound_usd']:.4f} דולר (מול {acc['recorded_cost_usd']:.4f} שנרשמו).", size=9.5, italic=True)
    on_status = on["status"]
    if on_status == "STOPPED_AT_CAP":
        para(doc, f"ON נעצר בתקרת הבקשות התפעולית הקפואה ({analysis['caps'][c.CONDITION_ON]} בקשות) לאחר {on['completed']} מתוך {on['planned']} מקרים; המקרים שלא הופקו נותרו בראיות כ-NOT_PRODUCED, ושלב 4 (חוקר השונות) לא הופעל. זהו תוצא שנרשם מראש במדיניות הכשלים הטכניים, לא כשל של המודל.", size=9.5, italic=True)
    para(doc, "טבלה מזווגת לפי מקרה (מספרי המקרים לפי מספור המסגרת המלאה):", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    table(doc, ["מקרה", "סטטוס ON", "סטטוס OFF", "שורות מיפוי ON/OFF", "לא מכוסים ON/OFF", "בקשות ON/OFF", "עלות ON/OFF ($)", "זמן ON~/OFF (ש׳)", "אפיזודות ON"], [
        [r["case_id"], r["on_status"], r["off_status"], f"{_f(r['on_mapping_rows'])}/{_f(r['off_mapping_rows'])}",
         f"{_f(r['on_uncovered'])}/{_f(r['off_uncovered'])}", f"{_f(r['on_requests'])}/{_f(r['off_requests'])}",
         f"{_f(r['on_cost_usd'], 4)}/{_f(r['off_cost_usd'], 4)}", f"{_s(r['on_elapsed_span_s'])}/{_s(r['off_elapsed_s'])}", _f(r["on_episodes"])]
        for r in analysis["paired"]
    ], size=7.5)
    para(doc, "הפרשים מזווגים (ON פחות OFF), ספירת מקרים לפי כיוון:", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    table(doc, ["מדד", "זוגות", "ON גבוה יותר", "OFF גבוה יותר", "שווה", "חציון הפרש", "ממוצע הפרש"], [
        [name, d["pairs"], d["on_higher"], d["off_higher"], d["equal"], _f(d["median_delta_on_minus_off"]), _f(d["mean_delta_on_minus_off"])]
        for name, d in (("שורות מיפוי", diff["mapping_rows"]), ("פרגמנטים לא מכוסים", diff["uncovered_fragments"]), ("בקשות", diff["requests"]), ("עלות (דולר)", diff["cost_usd"]), ("זמן (שניות)", diff["elapsed_seconds"]))
    ], size=8.5)
    para(doc, "ההפרשים מתארים ריצה זו בלבד; \"גבוה יותר\" אינו \"טוב יותר\".", size=9.5, italic=True)
    figure(doc, charts["C1"], _caption("C1", analysis))
    figure(doc, charts["C2"], _caption("C2", analysis))
    figure(doc, charts["C3"], _caption("C3", analysis))
    figure(doc, charts["C6"], _caption("C6", analysis))

    heading(doc, "5. תוצאות B — מדדי תקשורת (ON בלבד)", 1)
    routes = on["routes"].get("routes", [])
    para(doc, f"נרשמו {det['episodes_total']} אפיזודות, מהן {det['scientific_complete']} שלמות מדעית ({det['setting_level_episodes']} ברמת ההגדרה, השאר משויכות למקרים). סיבות סיום: {fmt_counts(det['termination_reasons'])}. מסלולי שאלות (מקור→יעד: מספר שאלות): " + (", ".join(f"{r['source_agent']}→{r['target_agent']}: {r['question_count']}" for r in routes) if routes else "אין") + f". מקרים ללא אפיזודה: {', '.join(det['cases_with_no_episode']) if det['cases_with_no_episode'] else 'אין'} — תצפית תקפה, לא כשל.")
    table(doc, ["מקרה", "אפיזודות", "שאלות", "תשובות", "בקשות משויכות", "אסימוני קלט / פלט", "תוויות Detector-v1"], [
        [r["case_id"], _f(r["on_episodes"]), _f(r["on_questions"]), _f(r["on_answers"]), _f(r["on_requests"]),
         f"{_f(r['on_prompt_tokens'])} / {_f(r['on_completion_tokens'])}", fmt_counts(r["on_detector_v1"], "NO_EPISODE")]
        for r in analysis["paired"]
    ], size=7.5)
    figure(doc, charts["C4"], _caption("C4", analysis))

    heading(doc, "6. תוצאות C — גבולות Detector-v1", 1)
    para(doc, "Detector-v1 הוחל לאחר הריצה על האפיזודות שנרשבו ב-ON כתווית \"מועמד לבדיקה אנושית\" בלבד. הוא אינו כותב תור אנושי, אינו משנה תשובה או מודל ואינו גורם לניסיון חוזר; ספיו לא שונו. תחת OFF אין יחידת ניתוח ולכן הוא לא רלוונטי — לא \"אפס התרעות\", והמכנה שלו לעולם אינו ממוזג עם מדדי OFF.")
    table(doc, ["סיווג", "אפיזודות"], [[k, v] for k, v in det["classifications"].items()] or [["אין אפיזודות", 0]], size=9, widths_cm=[6, 4])
    para(doc, f"אותות שנדלקו (אפיזודה יכולה להדליק כמה): {fmt_counts(det['reason_codes'])}. שיעור המקרים עם תווית מועמדת כלשהי: {_f(det['share_of_cases_with_any_candidate'], 3)}. תוויות לפי מקרה: " + "; ".join(f"{r['case_id']}: {fmt_counts(r['on_detector_v1'])}" for r in analysis["paired"]) + ".", size=10)
    figure(doc, charts["C5"], _caption("C5", analysis))

    heading(doc, "7. תוצאות D — מדדי רלוונטיות עסקית", 1)
    para(doc, f"עלות לפריט שהושלם: ON {_f(on['cost_per_completed_artifact_usd'], 6)} דולר, OFF {_f(off['cost_per_completed_artifact_usd'], 6)} דולר. בקשות לפריט שהושלם: ON {_f(on['requests_per_completed_artifact'], 2)}, OFF {_f(off['requests_per_completed_artifact'], 2)}. שיעור המקרים שבהם ON הפיק תווית מועמדת לבדיקה: {_f(det['share_of_cases_with_any_candidate'], 3)}. משמעות עסקית: ON קונה, במחיר של יותר בקשות ועלות, שובל תקשורת שניתן לבחון ולסנן; OFF זול ומהיר יותר אך אינו מייצר שובל כזה. אין להסיק עדיפות של תנאי על בסיס העלות בלבד, ואין להסיק שהתוויות נכונות.")

    heading(doc, "8. תוצאות E — הערכה אנושית", 1)
    para(doc, f"מצב: NOT_MEASURED. נוצרו {cards.get('cards', 0)} כרטיסים עיוורים ומוסווים (SHA-256 {str(cards.get('cards_sha256', ''))[:16]}…) בשורש הפרטי בלבד, ללא מזהה מקרה, אפיזודה, תנאי או תווית גלאי. שני מעריכים בלתי תלויים יחליטו לכל כרטיס: REVIEW_WORTHY / NOT_REVIEW_WORTHY / INSUFFICIENT_INFORMATION ויציעו פעולה: verify / clarify / revise guideline / no action. עד אז הסכמה, נכונות, דיוק, רגישות, F1, עומס ותועלת אינם נמדדים, ואין לסמן שיפוט של מודל כשיפוט אנושי.")
    figure(doc, charts["C7"], _caption("C7", analysis), width_cm=13)

    heading(doc, "9. מה נמצא", 1)
    for text in _found_he(analysis):
        bullet(doc, text)
    heading(doc, "10. מה לא נמצא", 1)
    for text in (
        "לא נמצא (ולא נבדק) איזה תנאי מפיק ניתוח טוב יותר: אין אמת-קרקע ואין ניקוד אנושי.",
        "לא נמצא האם תוויות Detector-v1 נכונות: אין תוויות אנושיות.",
        "לא נמצאה תועלת אנושית, חיסכון בעומס או שיפור בדיוק.",
        "לא נמצאה יציבות בין ריצות: ריצה אחת, יום אחד, טמפרטורת ברירת מחדל.",
    ):
        bullet(doc, text)
    heading(doc, "11. מה מותר לטעון", 1)
    for text in (
        f"על {n} מקרים מזווגים מקורפוס ציבורי אחד, תחת מדיניות בקשות זהה, שני הזרמים התנהגו כמתואר בסעיף 4 (שלמות, תקינות, בקשות, עלות, זמן, כשלים).",
        "ON מייצר ראיות תקשורת ניתנות לבחינה; OFF אינו מייצר כאלה מבנית.",
        "Detector-v1 הפיק תוויות מועמדות בכמות המדווחת; משמעותן תיקבע רק בהערכה אנושית.",
        "התוצאות הן ראיות אמפיריות פרוספקטיביות תיאוריות בלבד." if not fixture else "התוצאות הן תוצר הנדסי בלבד.",
    ):
        bullet(doc, text)
    heading(doc, "12. מה אסור לטעון", 1)
    for text in (
        "שהתרעות Detector-v1 נכונות.", "ש-VEGO-AI טוב יותר באופן כללי או באיכות.", "שתנאי עדיף בגלל עלות בלבד.",
        "ש-OFF הפיק אפס התרעות.", "שפלט של ריצת בדיקה הוא ראיה אמפירית.", "ששיפוט של מודל הוא שיפוט אנושי.",
        "שהתוצאות מוכללות מעבר לקורפוס, למודל, ליום ולתצורה.",
    ):
        bullet(doc, text)

    heading(doc, "13. פרשנות עסקית והצעד הבא", 1)
    para(doc, "מבחינה עסקית, השאלה אינה \"מי טוב יותר\" אלא \"מה קונה ההשקעה הנוספת\". ריצה זו מכמתת את ההשקעה (בקשות, עלות, זמן) ואת התוצר הנלווה שרק ON מייצר — שובל שאלה–תשובה מתועד שניתן לסנן בעזרת תווית דיווח. האם השובל הזה שווה את מחירו נקבע בשלב הבא: ניקוד עיוור של הכרטיסים על ידי שני מעריכים, חישוב הסכמה, ורק אז דיון בנכונות התוויות ובעומס העבודה שהן חוסכות או מוסיפות.")
    para(doc, f"הצעד הבא: שני מעריכים מנקדים {det['scientific_complete']} כרטיסים לפי המחוון בעברית (כ-30–45 דקות לכל מעריך); לאחר מכן חישוב הסכמה בין מעריכים ודיווח על התפלגות ההחלטות מול תוויות הגלאי, תחת גבול הטענות הקפוא.")

    heading(doc, "14. מגבלות", 1)
    for text in (
        "קורפוס ציבורי אחד שנוצר על ידי מודלי שפה; לא נתוני סטודנטים; לא ייצוגי.",
        f"{n} מקרים מזווגים שנדגמו בסיד מתוך מסגרת של {selection['eligible_case_count']}; שלושה חופפים בבייטים לקלטי מחקר 1 (ללא שימוש חוזר בפלטים).",
        "מודל ספק אחד ביום אחד בטמפרטורת ברירת המחדל; ריצה נוספת עשויה להיות שונה.",
        "עלות וזמן לכל מקרה ב-ON הם שיוכים תחת מקביליות, לא מדידות מבודדות; עלות רמת ההגדרה משותפת.",
        "מדדים מבניים מתארים פריטים ואינם שופטים נכונות; תוויות Detector-v1 הן מועמדות בלבד.",
        "ספירות שורות המיפוי אינן ניתנות להשוואה ישירה בין התנאים: ON ממפה מול הנחיות שיצר יועץ התחום, OFF מול הנחיות שגזר בעצמו מתיאור התחום, משום שההנחיה הקפואה של OFF אינה כוללת רשימת הנחיות משותפת.",
        "ON נעצר בתקרת הבקשות התפעולית לאחר תשעה מקרים ולפני שלב חוקר השונות; שלושה מקרים לא הופקו ב-ON, ולכן ההשוואה המזווגת המלאה מכסה תשעה זוגות.",
    ):
        bullet(doc, text)

    heading(doc, "נספח א — מפתח ראיות ושלמות", 1)
    table(doc, ["פריט", "SHA-256 / ערך"], [
        ["מניפסט קפוא", analysis["manifest_sha256"]],
        ["ראש Git של הביצוע", analysis["execution_git_sha"]],
        ["מקור Detector-v1", manifest["detector_v1"]["sha256"]],
        ["אורקסטרטור מוגן", manifest["conditions"][c.CONDITION_ON]["protected_runtime_sha256"]["orchestrator.py"]],
        ["תבנית ההנחיה של OFF", manifest["conditions"][c.CONDITION_OFF]["prompt_template_sha256"]],
        ["כרטיסי הערכה (פרטי)", str(cards.get("cards_sha256", "לא זמין"))],
        ["כלל האימות", "study2_prospective_run.py validate מחשב מחדש כל ערך מפורסם מהראיות הפרטיות"],
    ], size=8, widths_cm=[4.5, 12.5])
    heading(doc, "נספח ב — הגדרות המדדים הקפואות", 1)
    table(doc, ["מדד", "סעיף", "הגדרה", "מכנה", "סיווג ראיות"], [
        [m["id"], m["section"], m["definition"], m["denominator"], m["evidence_class"]] for m in manifest["metrics"]
    ], size=7.5, widths_cm=[1.2, 1.2, 8.2, 4.0, 2.4])
    heading(doc, "נספח ג — מילון מונחים", 1)
    table(doc, ["מונח", "משמעות במסמך זה"], [
        ["VEGO‑AI_ON", "זרם העבודה הרב-סוכני המוגן ללא שינוי: יועץ שפה, יועץ תחום, בוחן מודל, חוקר שונות, עם שאלה–תשובה מתועדת בין סוכנים"],
        ["VEGO‑AI_OFF", "זרם עבודה ישיר מותאם: אותו מודל, אותם קלטים, אותה סכמת פלט, בקשה אחת למקרה, ללא סוכנים, ללא שאלה–תשובה, ללא סבבים. אינו \"בלי בינה מלאכותית\""],
        ["אפיזודה", "רצף שאלות ותשובות בין סוכן מקור ליועץ, שנפתח בקריאה מפיקה ונסגר בהתכנסות או בהגעה למספר הסבבים המרבי"],
        ["Detector‑v1", "כלל קפוא המסווג אפיזודה שלמה ל-STRONG_ALERT (ביטחון נמוך, ראיה חסרה או סיום במספר סבבים מרבי), WEAK_ALERT (ביטחון בינוני או כמה סבבים) או NO_ALERT; תווית דיווח בלבד"],
        ["NOT_APPLICABLE", "אין יחידת ניתוח לגלאי (OFF); שונה מאפס התרעות ומ-NO_EPISODE"],
        ["NO_EPISODE", "מקרה תחת ON שלא יצר אפיזודה; תצפית תקפה, לא כשל"],
        ["NOT_MEASURED", "מדד שהוגדר אך לא נמדד, למשל כל מדד ההערכה האנושית עד שני מעריכים"],
        ["NOT_AVAILABLE", "מדד שאינו קיים מבנית בתנאי, למשל ראיות תקשורת תחת OFF"],
        ["פריט (artifact)", "פלט המקרה בחוזה המשותף study2-condition-output-v1: מיפוי הנחיות, סיכום כיסוי, פרגמנטים לא מכוסים"],
        ["בקשה", "קריאה אחת לספק, כולל ניסיונות תעבורה חוזרים וניסיונות פענוח חוזרים; כל בקשה נספרת בתקרות ובתקציב"],
        ["רזרבציה", "עלות מקסימלית מוקצית לבקשה לפני שליחתה: 12,000 אסימוני קלט ו-16,384 אסימוני פלט במחירון הקפוא"],
        ["שער", "בדיקה שנכשלת-סגור לפני הביצוע; כל שער שנכשל מסרב לבצע"],
        ["מזהה מקרה", "מספור המסגרת המלאה (01–21) של מודלי המועמד בקורפוס Text2UML AirTravel"],
    ], size=8.5, widths_cm=[3.5, 13.5])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def _found_he(analysis: dict[str, Any]) -> list[str]:
    cond = analysis["conditions"]
    on, off = cond[c.CONDITION_ON], cond[c.CONDITION_OFF]
    det = analysis["detector"]
    diff = analysis["differences"]
    on_rows, off_rows = on.get("mapping_rows_per_case") or [0], off.get("mapping_rows_per_case") or [0]
    on_unc, off_unc = on.get("uncovered_per_case") or [0], off.get("uncovered_per_case") or [0]
    guideline_count = analysis.get("on_reference_guideline_count")
    guideline_text = f" ({guideline_count} הנחיות ייחוס הופקו על ידי יועץ התחום)" if guideline_count is not None else ""
    found = [
        f"שני הזרמים הפיקו פריטים תקפים סכמטית ברוב המקרים: ON {on['completed']}/{on['planned']}, OFF {off['completed']}/{off['planned']}; שניהם יחד ב-{diff['both_completed']} מקרים.",
        f"מבנה הפריטים שונה מהותית: פריטי ON הכילו {min(on_rows)}–{max(on_rows)} שורות מיפוי לכל מקרה{guideline_text}, ואילו OFF, שאינו מקבל רשימת הנחיות, גזר בעצמו {min(off_rows)}–{max(off_rows)} הנחיות לכל מקרה. במקביל ON סימן {min(on_unc)}–{max(on_unc)} פרגמנטים לא מכוסים לכל מקרה (סה\"כ {on['uncovered_total']}: {fmt_counts(on.get('fragment_labels_total'))}) לעומת {min(off_unc)}–{max(off_unc)} ב-OFF (סה\"כ {off['uncovered_total']}: {fmt_counts(off.get('fragment_labels_total'))}). אלה פירוקים שונים של אותה משימה, לא מדד איכות.",
        f"Detector-v1 סימן {det['candidate_alerts']} מתוך {det['scientific_complete']} האפיזודות השלמות כמועמדות, רובן STRONG בגלל S1 (ביטחון נמוך של היועץ, {det['reason_codes'].get('S1_LOW_ANSWER_CONFIDENCE', 0)} אפיזודות); כאשר הכלל הקפוא מסמן כמעט כל אפיזודה, כושר ההבחנה שלו בריצה זו נמוך — תצפית תיאורית, לא טענה על נכונות.",
        f"ON דרש {on['requests']} בקשות לעומת {off['requests']} ב-OFF, ועלה {on['cost_usd']:.4f} דולר לעומת {off['cost_usd']:.4f}; סך הריצה {analysis['budget']['actual_cost_usd']:.4f} דולר, הרחק מתחת לתקרה.",
        f"ON יצר {det['episodes_total']} אפיזודות שאלה–תשובה מתועדות, ו-Detector-v1 סימן {det['candidate_alerts']} מהן כמועמדות לבדיקה ({fmt_counts(det['classifications'])}).",
        f"הפרשי מבנה: שורות מיפוי — ON גבוה ב-{diff['mapping_rows']['on_higher']} מקרים, OFF ב-{diff['mapping_rows']['off_higher']}, שווה ב-{diff['mapping_rows']['equal']}; פרגמנטים לא מכוסים — ON גבוה ב-{diff['uncovered_fragments']['on_higher']}, OFF ב-{diff['uncovered_fragments']['off_higher']}.",
    ]
    if on["technical_failures"] or off["technical_failures"]:
        found.append(f"כשלים טכניים נותרו בראיות ודווחו: ON {on['technical_failures']}, OFF {off['technical_failures']}.")
    return found


def build_executive_summary(analysis: dict[str, Any], out_path: Path) -> Path:
    doc = Document()
    setup(doc)
    cond = analysis["conditions"]
    on, off = cond[c.CONDITION_ON], cond[c.CONDITION_OFF]
    det = analysis["detector"]
    diff = analysis["differences"]
    n = len(analysis["case_ids"])
    heading(doc, "תקציר מנהלים — מחקר 2: ניסוי פרוספקטיבי מזווג VEGO-AI_ON לעומת VEGO-AI_OFF", 1)
    para(doc, f"ריצה {analysis['run_id']} · מודל {analysis['model']['model']} · ראש Git {analysis['execution_git_sha'][:12]} · סיווג ראיות: {_evidence_he(analysis)}", size=9.5, color=GREY, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(doc, f"מה נעשה: ריצה אחת מאושרת מראש של זרם העבודה הרב-סוכני (ON) מול זרם ישיר מותאם (OFF) על {n} מקרים שנדגמו בסיד מתוך 21 מודלים כשירים בקורפוס הציבורי Text2UML AirTravel, תחת תקרה קשיחה של 6.00 דולר ומדיניות בקשות זהה לשני התנאים.", size=10.5)
    heading(doc, "מה קרה", 2)
    for text in (
        f"שלמות: ON {on['completed']}/{on['planned']}, OFF {off['completed']}/{off['planned']}; שניהם השלימו {diff['both_completed']} מקרים.",
        f"תקינות מבנית: ON {on['summary_consistent_count']}/{on['completed']} עם סיכום כיסוי עקבי; OFF {off['summary_consistent_count']}/{off['completed']}.",
        f"בקשות: ON {on['requests']} (רמת הגדרה {on['setting_level_requests']}), OFF {off['requests']}; סך {analysis['budget']['requests']} מתוך תקרה {analysis['budget']['total_request_cap']}.",
        f"עלות: ON {on['cost_usd']:.4f}$, OFF {off['cost_usd']:.4f}$; סך נרשם {analysis['budget']['actual_cost_usd']:.4f}$ (חסם עליון {analysis['accounting']['spend_upper_bound_usd']:.4f}$ עם {analysis['accounting']['unrecorded_in_flight_requests']} בקשה בטיסה שחויבה ברזרבציה מלאה) מול רזרבציה {analysis['whole_study_reservation_usd']:.4f}$ ותקרה 6.00$.",
        f"זמן: ON {_s(on['elapsed_seconds'])} שניות, OFF {_s(off['elapsed_seconds'])} שניות לכל התנאי.",
        f"כשלים טכניים: ON {on['technical_failures']}, OFF {off['technical_failures']}; ניסיונות תעבורה חוזרים: ON {on['retried_requests']}, OFF {off['retried_requests']}.",
        f"ראיות תקשורת: ON {det['episodes_total']} אפיזודות ({det['scientific_complete']} שלמות); OFF — אין מבנית.",
        f"Detector-v1 (ON בלבד, דיווח בלבד): {fmt_counts(det['classifications'])}; שיעור מקרים עם תווית מועמדת {_f(det['share_of_cases_with_any_candidate'], 3)}. תחת OFF: לא רלוונטי.",
    ):
        bullet(doc, text, size=10)
    heading(doc, "מה זה מראה ומה לא", 2)
    para(doc, "מראה כיצד שני הזרמים התנהגו תפעולית על אותם קלטים תחת אותה מדיניות. אינו מראה מי הפיק ניתוח טוב יותר, האם תווית כלשהי נכונה, או תועלת אנושית — כל אלה NOT_MEASURED עד שני מעריכים.", size=10.5)
    heading(doc, "החלטה נדרשת מהמנחים", 2)
    para(doc, f"אישור לניקוד עיוור של {det['scientific_complete']} כרטיסים על ידי שני מעריכים (כ-30–45 דקות לכל אחד), לפי המחוון בעברית. רק לאחר מכן ניתן לחשב הסכמה ולדון בנכונות תוויות הגלאי.", size=10.5)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def build_cards_docx(cards_json: Path, out_path: Path) -> Path:
    cards = json.loads(cards_json.read_text(encoding="utf-8"))
    doc = Document()
    setup(doc)
    heading(doc, "מחקר 2 — כרטיסי הערכה עיוורת של אפיזודות שאלה–תשובה", 0)
    para(doc, "הנחיות: לכל כרטיס החליטו האם האפיזודה ראויה לבדיקה אנושית בעדיפות (REVIEW_WORTHY / NOT_REVIEW_WORTHY / INSUFFICIENT_INFORMATION), הציעו פעולה (verify / clarify / revise guideline / no action) ונקדו R1–R5 בסולם 1–5. אין לדון עם המעריך השני לפני סיום. אין מזהי מקרה, תנאי או תווית גלאי בכרטיסים. הטקסט המקורי באנגלית מוצג כפי שנוצר.", size=10)
    for index, card in enumerate(cards, start=1):
        heading(doc, f"כרטיס {index} — {card['card_id']}", 2)
        para(doc, f"סבבים: {card['round_count']} · סיום: {card['termination']}", size=9.5, color=GREY, align=WD_ALIGN_PARAGRAPH.RIGHT)
        for k, ex in enumerate(card["exchanges"], start=1):
            para(doc, f"חילוף {k} (יועץ {'שפה' if ex['target_advisor'] == 'language' else 'תחום'}) · ביטחון מדווח: {ex['reported_confidence'] or 'לא דווח'} · ראיה מדווחת: {'כן' if ex['reported_evidence'] else 'לא'}", size=9.5, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
            para(doc, "Q: " + ex["question"], size=9.5, rtl=False, align=WD_ALIGN_PARAGRAPH.LEFT)
            para(doc, "A: " + ex["answer"], size=9.5, rtl=False, align=WD_ALIGN_PARAGRAPH.LEFT)
            if ex["reported_evidence"]:
                para(doc, "Evidence: " + str(ex["reported_evidence"]), size=9, rtl=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=GREY)
        table(doc, ["החלטה", "פעולה", "R1", "R2", "R3", "R4", "R5", "הערה"], [["", "", "", "", "", "", "", ""]], size=9)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
