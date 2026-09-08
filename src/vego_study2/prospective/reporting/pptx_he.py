"""Hebrew right-to-left supervisor presentation (python-pptx, installed ad hoc)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from .. import constants as c
from .analysis import fmt_counts

HEB = "David"
_LATIN_TOKEN = re.compile(r"[A-Za-z][A-Za-z0-9_\-/]*(?:\.[A-Za-z0-9_\-/]+)*")


def _segments(text: str) -> list[tuple[str, bool]]:
    """Split into (segment, is_latin) so PowerPoint resolves each script itself.

    PowerPoint's DrawingML renderer scrambles digit-bearing Latin identifiers
    inside a Hebrew-tagged run; giving every Latin identifier its own en-US
    run and leaving all neutrals in the Hebrew runs renders correctly.
    """
    out: list[tuple[str, bool]] = []
    cursor = 0
    for match in _LATIN_TOKEN.finditer(text):
        if match.start() > cursor:
            out.append((text[cursor:match.start()], False))
        out.append((match.group(0), True))
        cursor = match.end()
    if cursor < len(text):
        out.append((text[cursor:], False))
    return out or [(text, False)]


def _format_run(run, size: float, bold: bool, color, latin: bool) -> None:
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = HEB
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "en-US" if latin else "he-IL")
    cs = rPr.find(qn("a:cs"))
    if cs is None:
        cs = rPr.makeelement(qn("a:cs"), {})
        rPr.append(cs)
    cs.set("typeface", HEB)


def _fill_paragraph(paragraph, text: str, size: float, bold: bool = False, color=None) -> None:
    from pptx.enum.text import PP_ALIGN

    paragraph.alignment = PP_ALIGN.RIGHT
    paragraph._p.get_or_add_pPr().set("rtl", "1")
    for segment, latin in _segments(text):
        run = paragraph.add_run()
        run.text = segment
        _format_run(run, size, bold, color, latin)


def _textbox(slide, left, top, width, height, lines: list[tuple[str, float, bool]], color: str | None = None):
    from pptx.util import Cm

    box = slide.shapes.add_textbox(Cm(left), Cm(top), Cm(width), Cm(height))
    frame = box.text_frame
    frame.word_wrap = True
    first = True
    for text, size, bold in lines:
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        _fill_paragraph(paragraph, text, size, bold, color)
    return box


def _title(slide, text: str, subtitle: str | None = None) -> None:
    _textbox(slide, 1.0, 0.6, 31.3, 2.0, [(text, 26, True)], color="1F3A5F")
    if subtitle:
        _textbox(slide, 1.0, 2.5, 31.3, 1.2, [(subtitle, 13, False)], color="595959")


def _footer(slide, analysis: dict[str, Any], text: str) -> None:
    _textbox(slide, 1.0, 17.6, 31.3, 0.9, [(f"{text} · ריצה {analysis['run_id']} · סיווג ראיות: {_evidence_he(analysis)}", 9, False)], color="7F7F7F")


def _evidence_he(analysis: dict[str, Any]) -> str:
    return {"PROSPECTIVE EMPIRICAL EVIDENCE": "ראיות אמפיריות פרוספקטיביות",
            "ENGINEERING-ONLY FIXTURE": "תוצר הנדסי בלבד (ספק מדומה)"}.get(analysis["evidence_class"], analysis["evidence_class"])


def _s(value) -> str:
    return "לא זמין" if value is None else f"{value:.0f}"


def build_presentation(analysis: dict[str, Any], manifest: dict[str, Any], selection: dict[str, Any], charts: dict[str, Path], out_path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Cm

    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    blank = prs.slide_layouts[6]
    cond = analysis["conditions"]
    on, off = cond[c.CONDITION_ON], cond[c.CONDITION_OFF]
    det = analysis["detector"]
    diff = analysis["differences"]
    n = len(analysis["case_ids"])
    fixture = analysis["mode"] != "LIVE"

    s = prs.slides.add_slide(blank)
    _textbox(s, 1.0, 5.5, 31.3, 3, [("מחקר 2 — ניסוי פרוספקטיבי מזווג", 40, True)], color="1F3A5F")
    _textbox(s, 1.0, 8.6, 31.3, 2, [("VEGO-AI_ON לעומת VEGO-AI_OFF: השוואה תפעולית תיאורית על קורפוס ציבורי אחד", 20, False)], color="404040")
    _textbox(s, 1.0, 11.5, 31.3, 3, [
        (f"ריצה {analysis['run_id']} · מודל {analysis['model']['model']} · ראש Git {analysis['execution_git_sha'][:12]}", 14, False),
        (f"סיווג ראיות: {_evidence_he(analysis)} · הערכה אנושית: NOT_MEASURED", 14, True),
        ("מצגת למנחים — איריס וארנון", 14, False),
    ], color="595959")
    if fixture:
        _textbox(s, 1.0, 15.5, 31.3, 1.5, [("אזהרה: ריצת בדיקה עם ספק מדומה — המספרים הם תוצר הנדסי בלבד", 16, True)], color="A6291F")

    s = prs.slides.add_slide(blank)
    _title(s, "השאלה וגבול הטענות")
    _textbox(s, 1.0, 4.0, 31.3, 12, [
        ("שאלה ראשית (שני התנאים): כשהמודל, המקרים, חומר התחום, סכמת הפלט, תקרת הפלט ומדיניות הבקשות קבועים — כיצד נבדלים הזרם הרב-סוכני והזרם הישיר בשלמות, בתקינות, בעלות, בזמן, בכשלים ובזמינות ראיות תקשורת?", 16, False),
        ("שאלה משנית (ON בלבד): האם Detector-v1 מזהה אפיזודות הראויות לבדיקה אנושית בעדיפות? — NOT_MEASURED עד ניקוד אנושי.", 16, False),
        ("", 8, False),
        ("ניסוי חקרני מזווג. לא טענת עדיפות, לא תועלת אנושית, לא נכונות התרעות, לא דיוק/רגישות/F1, לא סיבתיות, לא ייצוגיות, לא הכללה.", 16, True),
        ("OFF אינו \"בלי בינה מלאכותית\": אותו מודל, אותם קלטים, אותה סכמה — בלי סוכנים, בלי שאלה–תשובה, בלי לולאת סבבים.", 15, False),
    ])
    _footer(s, analysis, "שקף 2")

    s = prs.slides.add_slide(blank)
    _title(s, "עיצוב הניסוי: מה זהה ומה שונה")
    _textbox(s, 17.5, 4.0, 14.8, 12, [
        ("VEGO-AI_ON", 18, True),
        ("ארבעה סוכנים: יועץ שפה, יועץ תחום, בוחן מודל, חוקר שונות", 14, False),
        ("שאלה–תשובה בין סוכנים, מוגבלת ומתועדת", 14, False),
        ("Detector-v1: תווית דיווח בלבד על האפיזודות", 14, False),
        (f"תקרה תפעולית: {manifest['caps']['on_requests_per_case']} בקשות למקרה", 14, False),
    ], color="1F4E79")
    _textbox(s, 1.0, 4.0, 15.5, 12, [
        ("VEGO-AI_OFF", 18, True),
        ("קריאה ישירה אחת לכל מקרה", 14, False),
        ("ללא סוכנים, ללא שאלה–תשובה, ללא סבבים", 14, False),
        ("Detector-v1: לא רלוונטי (אין אפיזודות; לא \"אפס\")", 14, False),
        (f"תקרה תפעולית: {manifest['caps']['off_requests_per_case']} בקשות למקרה", 14, False),
    ], color="C55A11")
    _textbox(s, 1.0, 12.5, 31.3, 4.5, [
        ("משותף לשניים: gpt-5.6-luna · max_completion_tokens 16384 · טמפרטורה לא נשלחת (ברירת מחדל) · זמן קצוב 180 ש׳ · ניסיון תעבורה חוזר אחד · אותו לקוח מוגן · אותה סכמת פלט · אותו שומר תקציב", 14, True),
        ("סדר: OFF ואז ON · מקביליות 2 · יעד רשת יחיד api.openai.com · אישור נקרא רק על ידי ה-SDK", 13, False),
    ], color="404040")
    _footer(s, analysis, "שקף 3")

    s = prs.slides.add_slide(blank)
    _title(s, "נתונים, בחירת מקרים ותקציב")
    _textbox(s, 1.0, 4.0, 31.3, 13, [
        (f"קורפוס: Text2UML AirTravel (ציבורי, נוצר על ידי מודלי שפה, לא ייצוגי) · מסגרת כשירה: {selection['eligible_case_count']} מודלים · כל תקציר SHA-256 אומת", 15, False),
        (f"בחירה: סיד {selection['selection_seed']}, {n} מקרים — {', '.join(selection['selected_case_ids'])}; הוצאו {len(selection['excluded_case_ids'])}; חפיפה בבייטים למחקר 1: {len(selection['study1_overlap_by_digest'])} קבצים (ללא שימוש חוזר בפלט)", 15, False),
        ("", 6, False),
        (f"תקציב: תקרה 6.00$ · שומר 5.90$ · רזרבציה לבקשה {analysis['budget']['per_request_reserve_usd']:.5f}$ (12k קלט + 16,384 פלט) · תקרת בקשות {analysis['budget']['total_request_cap']}", 15, False),
        (f"תפריט: A — 12 מקרים (רזרבציה {analysis['whole_study_reservation_usd']:.2f}$) · B — 6×2 חזרות (זהה) · C — 4 מקרים (1.94$) → נבחר A: יחידות מזווגות נפרדות חשובות יותר מחזרות", 15, False),
        ("", 6, False),
        (f"בפועל: {analysis['budget']['requests']} בקשות · {analysis['budget']['actual_cost_usd']:.4f}$ · {analysis['budget']['prompt_tokens']} אסימוני קלט · {analysis['budget']['completion_tokens']} אסימוני פלט · חריגות שומר: {len(analysis['budget']['refusals'])}", 15, True),
    ])
    _footer(s, analysis, "שקף 4")

    for key, title, bullets, footer in (
        ("C1", "תוצאות A — שלמות פלט ותקינות מבנית", [
            f"ON {on['completed']}/{on['planned']} · OFF {off['completed']}/{off['planned']} · שניהם {diff['both_completed']} מקרים",
            f"סיכום כיסוי עקבי: ON {on['summary_consistent_count']}/{on['completed']} · OFF {off['summary_consistent_count']}/{off['completed']}",
            f"כשלים טכניים: ON {on['technical_failures']} · OFF {off['technical_failures']}",
        ], "שקף 5"),
        ("C2", "תוצאות A — בקשות ועלות", [
            f"ON {on['requests']} בקשות ({on['setting_level_requests']} ברמת ההגדרה) · {on['cost_usd']:.4f}$",
            f"OFF {off['requests']} בקשות · {off['cost_usd']:.4f}$",
            f"עלות לפריט שהושלם: ON {on['cost_per_completed_artifact_usd'] if on['cost_per_completed_artifact_usd'] is not None else 'לא זמין'} · OFF {off['cost_per_completed_artifact_usd'] if off['cost_per_completed_artifact_usd'] is not None else 'לא זמין'}",
            "עלות מתארת את הריצה; אינה טענת עדיפות",
        ], "שקף 6"),
        ("C3", "תוצאות A — זמן", [
            f"ON {_s(on['elapsed_seconds'])} שניות · OFF {_s(off['elapsed_seconds'])} שניות לכל התנאי",
            f"הפרש זמן לכל מקרה: ON גבוה ב-{diff['elapsed_seconds']['on_higher']} מקרים, OFF ב-{diff['elapsed_seconds']['off_higher']}",
            "טווחי ON לכל מקרה מקורבים תחת מקביליות 2",
        ], "שקף 7"),
        ("C5", "תוצאות B–C — תקשורת ו-Detector-v1 (ON בלבד)", [
            f"{det['episodes_total']} אפיזודות ({det['scientific_complete']} שלמות) · סיום: {fmt_counts(det['termination_reasons'])}",
            f"סיווגים: {fmt_counts(det['classifications'])} · אותות: {fmt_counts(det['reason_codes'])}",
            f"מקרים ללא אפיזודה: {', '.join(det['cases_with_no_episode']) or 'אין'} (תצפית תקפה) · OFF: לא רלוונטי",
            "תווית = מועמדת לבדיקה, לא בעיה מאומתת",
        ], "שקף 8"),
    ):
        s = prs.slides.add_slide(blank)
        _title(s, title)
        s.shapes.add_picture(str(charts[key]), Cm(0.8), Cm(4.0), width=Cm(19.5))
        _textbox(s, 21.0, 4.2, 12.0, 12, [(b, 14, False) for b in bullets], color="333333")
        _footer(s, analysis, footer)

    s = prs.slides.add_slide(blank)
    _title(s, "מה נמצא · מה לא נמצא · מה מותר · מה אסור")
    _textbox(s, 17.5, 4.0, 15.5, 6.5, [("נמצא", 16, True),
                                       (f"שני הזרמים הפיקו פריטים תקפים ברוב המקרים ({diff['both_completed']} יחד)", 12.5, False),
                                       (f"ON: {on['requests']} בקשות, {on['cost_usd']:.4f}$; OFF: {off['requests']}, {off['cost_usd']:.4f}$", 12.5, False),
                                       (f"ON יצר {det['episodes_total']} אפיזודות; {det['candidate_alerts']} מועמדות לבדיקה", 12.5, False)], color="1F4E79")
    _textbox(s, 1.0, 4.0, 15.5, 6.5, [("לא נמצא", 16, True),
                                      ("איזה זרם מפיק ניתוח טוב יותר (אין אמת-קרקע, אין ניקוד)", 12.5, False),
                                      ("האם תוויות הגלאי נכונות (אין תוויות אנושיות)", 12.5, False),
                                      ("תועלת אנושית, חיסכון בעומס, יציבות בין ריצות", 12.5, False)], color="A6291F")
    _textbox(s, 17.5, 10.8, 15.5, 6.5, [("מותר לטעון", 16, True),
                                        ("התנהגות תפעולית תיאורית של שני הזרמים על 12 מקרים מזווגים", 12.5, False),
                                        ("ON מייצר ראיות תקשורת; OFF לא — מבנית", 12.5, False),
                                        ("Detector-v1 הפיק תוויות מועמדות בכמות המדווחת", 12.5, False)], color="548235")
    _textbox(s, 1.0, 10.8, 15.5, 6.5, [("אסור לטעון", 16, True),
                                       ("שהתרעות נכונות · שVEGO-AI טוב יותר · שעלות = עדיפות", 12.5, False),
                                       ("ש-OFF הפיק אפס התרעות · ששיפוט מודל = שיפוט אנושי", 12.5, False),
                                       ("שהתוצאות מוכללות מעבר לקורפוס, למודל וליום", 12.5, False)], color="7F7F7F")
    _footer(s, analysis, "שקף 9")

    s = prs.slides.add_slide(blank)
    _title(s, "הצעד הבא והחלטות למנחים")
    _textbox(s, 1.0, 4.0, 31.3, 12, [
        (f"1. ניקוד עיוור: שני מעריכים, {det['scientific_complete']} כרטיסים, מחוון בעברית (REVIEW_WORTHY / NOT_REVIEW_WORTHY / INSUFFICIENT_INFORMATION; verify / clarify / revise guideline / no action; R1–R5). כ-30–45 דקות לכל מעריך.", 15, False),
        ("2. לאחר הניקוד: הסכמה בין מעריכים, התפלגות ההחלטות מול תוויות הגלאי — תחת גבול הטענות הקפוא. עד אז: NOT_MEASURED.", 15, False),
        ("3. החלטה נדרשת: מי שני המעריכים, ומועד.", 15, True),
        ("", 6, False),
        ("שקיפות: הריצה הוקפאה לפני כל פלט (פרה-רגיסטרציה, מניפסט חתום-עצמית, בחירת מקרים בסיד, רזרבציית תקציב); כל ערך מפורסם ניתן לחישוב מחדש מהראיות הפרטיות; אין בייטים של קורפוס, הנחיות או תשובות ב-Git.", 13, False),
    ])
    _footer(s, analysis, "שקף 10")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path
