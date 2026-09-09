"""Hebrew right-to-left slide primitives for python-pptx.

Two things are easy to get wrong when a PowerPoint deck is written in Hebrew, and both are
invisible until the file is opened on someone else's machine:

  1. Paragraph direction. `PP_ALIGN.RIGHT` only moves the text box content; it does not set the
     paragraph's `rtl` flag, so punctuation, parentheses and mixed Hebrew/Latin runs still lay
     out left-to-right. The flag lives on `a:pPr/@rtl` and has to be written directly.
  2. The complex-script typeface. `font.name` sets only `a:latin`. Hebrew is a complex script,
     so PowerPoint reads `a:cs`; leaving it unset makes PowerPoint substitute a fallback font
     for every Hebrew glyph while Latin runs keep the requested font.

A third defect only shows up once real content arrives: a Latin or numeric token inside a Hebrew
sentence takes its direction from its neighbours, so "44 וטקסט" renders glued and a trailing
colon after a Latin path jumps to the far end of the line. Unicode isolates (FSI/PDI) are the
textbook remedy, but PowerPoint's PDF export has no glyph for them and prints a missing-glyph
box, so they are deliberately NOT used here. Instead Latin and numeric tokens are kept out of
Hebrew prose by construction - they live in stat cards, table cells, bar labels, or a dedicated
`ltr_para` - which is also the better slide design.
"""

from __future__ import annotations

import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x14, 0x37, 0x5E)
ACCENT = RGBColor(0x2A, 0x78, 0xD6)
ACCENT_L = RGBColor(0xE2, 0xEE, 0xFD)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
ORANGE_L = RGBColor(0xFD, 0xEC, 0xE5)
GREEN = RGBColor(0x3C, 0x96, 0x5A)
GREEN_L = RGBColor(0xE6, 0xF4, 0xE9)
PURPLE = RGBColor(0x30, 0x1A, 0x6E)
PURPLE_L = RGBColor(0xE8, 0xE2, 0xFD)
GREY = RGBColor(0xAA, 0xAA, 0xA6)
GREY_L = RGBColor(0xF0, 0xF0, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x0B, 0x0B, 0x0B)
MUTED = RGBColor(0x52, 0x51, 0x4E)
LINE = RGBColor(0xD8, 0xD8, 0xD4)
SOFT = RGBColor(0xF4, 0xF5, 0xF3)

FONT = "Segoe UI"
MONO = "Consolas"

W, H = Inches(13.333), Inches(7.5)
M = Inches(0.55)
CONTENT_W = W - 2 * M


HEBREW = re.compile(r"[֐-׿]")


def has_mixed_direction(text: str) -> bool:
    """True when Hebrew and a Latin/numeric token share one string, which lays out unreliably."""
    return bool(HEBREW.search(text)) and bool(re.search(r"[A-Za-z0-9]", text))


def emu(value) -> Emu:
    """Geometry must be integral EMU: dividing an Emu by an int yields a float, and python-pptx
    writes that float straight into the XML, which PowerPoint then refuses to open."""
    return Emu(int(value))


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    return prs


def slide(prs, *, fill=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = fill
    return s


def _set_font(run, *, size, bold, italic, color, font):
    f = run.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = color
    # a:latin is written by font.name; Hebrew needs a:cs or PowerPoint substitutes a fallback.
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:cs", "a:ea"):
        existing = rPr.find(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag[2:]}")
        if existing is None:
            existing = rPr.makeelement(
                f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag[2:]}", {})
            rPr.append(existing)
        existing.set("typeface", font)


def _rtl(paragraph, align):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("rtl", "1")
    paragraph.alignment = align


def tb(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box, tf


def para(tf, text, *, size=14, bold=False, italic=False, color=INK, font=FONT,
         align=PP_ALIGN.RIGHT, space_after=0, space_before=0, line=None, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    _rtl(p, align)
    p.space_after, p.space_before = Pt(space_after), Pt(space_before)
    if line:
        p.line_spacing = line
    run = p.add_run()
    run.text = text
    _set_font(run, size=size, bold=bold, italic=italic, color=color, font=font)
    return p


def ltr_para(tf, text, *, size=13, bold=False, color=INK, font=MONO,
             align=PP_ALIGN.CENTER, space_after=0, space_before=0, first=False):
    """A left-to-right line for code, formulas and paths, which are not Hebrew text."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pPr.set("rtl", "0")
    p.alignment = align
    p.space_after, p.space_before = Pt(space_after), Pt(space_before)
    run = p.add_run()
    run.text = text
    _set_font(run, size=size, bold=bold, italic=False, color=color, font=font)
    return p


def rich(tf, chunks, *, size=14, align=PP_ALIGN.RIGHT, space_after=0, space_before=0,
         line=None, first=False):
    """One paragraph built from differently styled runs: [(text, {bold:True, ...}), ...]."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    _rtl(p, align)
    p.space_after, p.space_before = Pt(space_after), Pt(space_before)
    if line:
        p.line_spacing = line
    for text, st in chunks:
        run = p.add_run()
        run.text = text
        _set_font(run, size=st.get("size", size), bold=st.get("bold", False),
                  italic=st.get("italic", False), color=st.get("color", INK),
                  font=st.get("font", FONT))
    return p


def bullets(tf, items, *, size=13, color=INK, space_after=5, bullet="• "):
    for text in items:
        para(tf, f"{bullet}{text}", size=size, color=color, space_after=space_after)


def card(slide, x, y, w, h, *, fill=SOFT, line_col=None, dashed=False, radius=0.05):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(w), emu(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line_col is not None:
        sh.line.color.rgb = line_col
        sh.line.width = Pt(1.25)
        if dashed:
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.07)
    return sh


def title(slide, text, *, sub=None, y=None):
    y = y if y is not None else M
    box, tf = tb(slide, M, y, CONTENT_W, Inches(0.72))
    para(tf, text, size=27, bold=True, color=NAVY, first=True)
    if sub:
        para(tf, sub, size=12.5, color=MUTED, space_before=3)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, M, y + Inches(0.78 if sub else 0.52),
                                  CONTENT_W, Pt(2))
    rule.fill.solid()
    rule.fill.fore_color.rgb = ACCENT
    rule.line.fill.background()
    rule.shadow.inherit = False
    return box


def chip(slide, x, y, text, *, kind="arch", w=Inches(3.15), h=Inches(0.29)):
    palette = {"arch": (ACCENT_L, NAVY, ACCENT), "fix": (ORANGE_L, ORANGE, ORANGE),
               "pub": (GREEN_L, GREEN, GREEN), "na": (GREY_L, MUTED, GREY),
               "pro": (PURPLE_L, PURPLE, PURPLE)}
    fill, ink, edge = palette[kind]
    sh = card(slide, x, y, w, h, fill=fill, line_col=edge, radius=0.35)
    tf = sh.text_frame
    para(tf, text, size=9.5, bold=True, color=ink, align=PP_ALIGN.CENTER, first=True)
    return sh


def stat(slide, x, y, w, h, value, label, *, detail=None, accent=NAVY, fill=SOFT):
    sh = card(slide, x, y, w, h, fill=fill, line_col=LINE)
    tf = sh.text_frame
    para(tf, value, size=30, bold=True, color=accent, align=PP_ALIGN.CENTER, first=True)
    para(tf, label, size=11, color=INK, align=PP_ALIGN.CENTER, space_before=2)
    if detail:
        para(tf, detail, size=8.5, color=MUTED, align=PP_ALIGN.CENTER, space_before=1)
    return sh


def bar_row(slide, x, y, w, label, value, top, shown, *, fixture=False,
            label_w=Inches(2.5), h=Inches(0.26)):
    """One right-anchored horizontal bar: the track grows leftwards, as Hebrew reads."""
    box, tf = tb(slide, x + w - label_w, y - Inches(0.02), label_w, h)
    para(tf, label, size=10.5, color=INK, first=True)
    track_w = w - label_w - Inches(0.12)
    track = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x), emu(y), emu(track_w), emu(h))
    track.adjustments[0] = 0.3
    track.fill.solid()
    track.fill.fore_color.rgb = GREY_L
    track.line.fill.background()
    track.shadow.inherit = False
    frac = (value / top) if top else 0
    bar_w = max(int(track_w * frac), Inches(0.05) if value == 0 else Inches(0.12))
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, emu(x + track_w - bar_w), emu(y),
                                 emu(bar_w), emu(h))
    bar.adjustments[0] = 0.3
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREY if value == 0 else (ORANGE if fixture else ACCENT)
    bar.line.fill.background()
    bar.shadow.inherit = False
    vbox, vtf = tb(slide, x + Inches(0.06), y - Inches(0.01), Inches(1.0), h)
    para(vtf, shown, size=9.5, bold=True, color=NAVY, align=PP_ALIGN.LEFT, first=True)
    return track


def rtl_table(slide, x, y, w, h, headers, rows, *, size=9.5, header_size=9.5,
              col_widths=None, highlight=None):
    """A table PowerPoint lays out right-to-left, so column 0 is the rightmost column."""
    highlight = highlight or set()
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), emu(x), emu(y), emu(w), emu(h))
    tbl = shape.table
    tbl._tbl.tblPr.set("rtl", "1")
    tbl.first_row = True
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = emu(cw)
    for c, text in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_L
        cell.margin_left = cell.margin_right = Inches(0.05)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        para(cell.text_frame, text, size=header_size, bold=True, color=NAVY,
             align=PP_ALIGN.CENTER, first=True)
    for r, cells in enumerate(rows, start=1):
        for c, text in enumerate(cells):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xFF, 0xF8, 0xE8) if (r - 1) in highlight else WHITE
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            mono = str(text).replace("_", "").isascii() and any(ch.isalpha() for ch in str(text))
            para(cell.text_frame, str(text), size=size, color=INK,
                 font=MONO if mono else FONT,
                 align=PP_ALIGN.CENTER if c else PP_ALIGN.RIGHT,
                 bold=(r - 1) in highlight, first=True)
    return tbl


def footer(slide, text, *, kind="def"):
    box, tf = tb(slide, M, H - M - Inches(0.3), CONTENT_W, Inches(0.3))
    colour = {"arch": NAVY, "fix": ORANGE, "def": MUTED, "na": MUTED}[kind]
    para(tf, text, size=8.5, color=colour, first=True)
    return box
